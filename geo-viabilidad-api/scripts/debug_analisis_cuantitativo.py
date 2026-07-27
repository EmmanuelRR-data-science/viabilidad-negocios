#!/usr/bin/env python
"""Dump del motor analítico cuantitativo sin FODA/LLM — JSON + Excel opcional.

Uso:
    cd geo-viabilidad-api
    uv run python scripts/debug_analisis_cuantitativo.py \
        --lat 19.4326 --lng -99.1332 --radio 1000 --rubro cafeteria --excel

Genera:
    scratch/debug_cuantitativo_<timestamp>.json
    scratch/debug_cuantitativo_<timestamp>.xlsx  (si --excel)
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

if sys.platform == "win32":
    os.environ.setdefault("PGPASSFILE", r"C:\Users\Public\pgpass.conf")
os.environ.setdefault("DEV_MODE", "True")


def main() -> None:
    parser = argparse.ArgumentParser(description="Debug cuantitativo sin IA")
    parser.add_argument("--lat", type=float, required=True)
    parser.add_argument("--lng", type=float, required=True)
    parser.add_argument("--radio", type=int, default=1000)
    parser.add_argument("--rubro", type=str, default="cafeteria")
    parser.add_argument("--tier", type=str, default="premium")
    parser.add_argument("--excel", action="store_true", help="Generar archivo Excel además de JSON")
    parser.add_argument("--out-dir", type=str, default="scratch")
    args = parser.parse_args()

    from app.clients.v0.database import SessionLocal
    from app.services.v0.analytics.analytics_service import procesar_calculo_analitico

    db = SessionLocal()
    try:
        resultado = procesar_calculo_analitico(
            db=db,
            lat=args.lat,
            lng=args.lng,
            radio=args.radio,
            rubro=args.rubro,
            tier=args.tier,
        )
    finally:
        db.close()

    os.makedirs(args.out_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = f"debug_cuantitativo_{ts}"

    json_path = os.path.join(args.out_dir, f"{base_name}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(resultado, f, ensure_ascii=False, indent=2, default=str)
    print(f"[OK] JSON → {json_path}")

    if args.excel:
        xlsx_path = os.path.join(args.out_dir, f"{base_name}.xlsx")
        _escribir_excel(resultado, xlsx_path)
        print(f"[OK] Excel → {xlsx_path}")


def _escribir_excel(data: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    # --- Sheet 1: KPIs ---
    ws = wb.active
    ws.title = "KPIs"
    kpis = [
        ("rubro", data.get("rubro")),
        ("poblacion_ponderada", data.get("poblacion_ponderada")),
        ("densidad_hab_km2", data.get("densidad_hab_km2")),
        ("competidores_conteo", data.get("competidores_conteo")),
        ("competidores_activos_conteo", data.get("competidores_activos_conteo")),
        ("distancia_competidor_cercano", data.get("distancia_competidor_cercano")),
        ("isc", data.get("isc")),
        ("score_demog", data.get("score_demog")),
        ("score_competencia", data.get("score_competencia")),
        ("score_trafico", data.get("score_trafico")),
        ("sva", data.get("sva")),
        ("bancos_conteo", data.get("bancos_conteo")),
        ("escuelas_conteo", data.get("escuelas_conteo")),
        ("transporte_conteo", data.get("transporte_conteo")),
    ]
    ws.append(["Métrica", "Valor"])
    ws["A1"].font = bold
    ws["B1"].font = bold
    for k, v in kpis:
        ws.append([k, v])

    # --- Sheet 2: NSE ---
    ws_nse = wb.create_sheet("NSE")
    nse = data.get("nse") or {}
    ws_nse.append(["Campo", "Valor"])
    ws_nse["A1"].font = bold
    ws_nse["B1"].font = bold
    for k, v in nse.items():
        if isinstance(v, dict):
            for sk, sv in v.items():
                ws_nse.append([f"{k}.{sk}", sv])
        else:
            ws_nse.append([k, v])

    # --- Sheet 3: Competidores ---
    ws_comp = wb.create_sheet("Competidores")
    comps = data.get("competidores_listado") or []
    comp_cols = ["nombre", "rating", "user_ratings_total", "distancia_metros", "place_id", "direccion"]
    ws_comp.append(comp_cols)
    for i, _h in enumerate(comp_cols, 1):
        ws_comp.cell(1, i).font = bold
    for c in comps:
        ws_comp.append([c.get(col) for col in comp_cols])

    # --- Sheet 4: Aliados ---
    ws_ally = wb.create_sheet("Aliados")
    aliados = data.get("aliados_listado") or []
    ally_cols = ["nombre", "tipo", "distancia_metros", "rating", "place_id"]
    ws_ally.append(ally_cols)
    for i, _h in enumerate(ally_cols, 1):
        ws_ally.cell(1, i).font = bold
    for a in aliados:
        ws_ally.append([a.get(col) for col in ally_cols])

    # --- Sheet 5: Afluencia ---
    ws_afl = wb.create_sheet("Afluencia")
    afl = data.get("afluencia_peatonal") or {}
    ws_afl.append(["Campo", "Valor"])
    ws_afl["A1"].font = bold
    ws_afl["B1"].font = bold
    for k in ("status", "venue_name", "dia_pico", "hora_pico", "saturación_promedio"):
        ws_afl.append([k, afl.get(k)])

    semanal = afl.get("afluencia_semanal") or {}
    if semanal:
        ws_afl_sem = wb.create_sheet("Afluencia_Semanal")
        header = ["Día"] + [f"{h:02d}:00" for h in range(24)]
        ws_afl_sem.append(header)
        for i, _h in enumerate(header):
            ws_afl_sem.cell(1, i + 1).font = bold
        for dia, curva in semanal.items():
            if isinstance(curva, list):
                ws_afl_sem.append([dia] + curva)

    # --- Sheet 6: Segmentación ---
    ws_seg = wb.create_sheet("Segmentacion")
    seg = data.get("segmentacion_demografica") or {}
    ws_seg.append(["Campo", "Valor"])
    ws_seg["A1"].font = bold
    ws_seg["B1"].font = bold
    for k, v in seg.items():
        if isinstance(v, (dict, list)):
            ws_seg.append([k, json.dumps(v, ensure_ascii=False, default=str)])
        else:
            ws_seg.append([k, v])

    for ws_item in wb.worksheets:
        for col in ws_item.columns:
            max_len = 0
            for cell in col:
                try:
                    val_len = len(str(cell.value or ""))
                    if val_len > max_len:
                        max_len = val_len
                except Exception:
                    pass
            ws_item.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    wb.save(path)


if __name__ == "__main__":
    main()
